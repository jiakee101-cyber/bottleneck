"""Build NAV_Bottleneck_Data_Template_v4_5.xlsx from v4_4 plus the Tabular Editor script.

Usage: python3 -I build.py <in.xlsx> <out.xlsx> <out.csx>
Only APPENDS columns/sheets/rows, except the three agreed fixes (marked FIX below).
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter, column_index_from_string as cidx

from measures import MEASURES, VISUALS, RELATIONSHIPS, SORT_BY, PQ_NOTE

src, dst, csx = sys.argv[1:4]
wb = openpyxl.load_workbook(src)

ARIAL = "Arial"
GREEN = "FF1F6F5C"   # input header
NAVY = "FF123A5E"    # derived header
BLUE_TXT = "FF0000FF"
GREY_TXT = "FF6B7684"
F_NUM = '#,##0;\\(#,##0\\);\\-'
F_DT = 'yyyy\\-mm\\-dd\\ hh:mm'


def hdr(cell, text, fill=NAVY, wrap=True):
    cell.value = text
    cell.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFFFF")
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(wrap_text=wrap, vertical="center")


def body(cell, value, fmt="General", color=None, wrap=False):
    cell.value = value
    cell.font = Font(name=ARIAL, size=10, color=color)
    cell.number_format = fmt
    if wrap:
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def add_columns(ws, start_col, specs, last_row, note_prefix):
    """specs: list of (header, formula_template_with_{r}, number_format, width, note)."""
    c0 = cidx(start_col)
    for i, (name, f, fmt, width, note) in enumerate(specs):
        col = c0 + i
        L = get_column_letter(col)
        assert ws.cell(1, col).value is None, f"{ws.title}!{L}1 not empty"
        hdr(ws.cell(1, col), name)
        ws.cell(1, col).comment = Comment(f"{note_prefix}: {note}", "v4.5")
        for r in range(2, last_row + 1):
            body(ws.cell(r, col), f.format(r=r), fmt)
        ws.column_dimensions[L].width = width


# ---------------------------------------------------------------- Constants
ws = wb["Constants"]
assert ws["A17"].value is None
body(ws["A17"], "Thin slack threshold (mins)")
body(ws["B17"], 0, "0", BLUE_TXT)
body(ws["C17"], "INPUT — v4.5. A custody, TA or broker feed is THIN when its slack to the pricing cut-off "
     "is at or below this. 0 = zero tolerance: thin only when the feed's cut-off is at or after the pricing "
     "cut-off, so any lateness reaches NAV directly. Separate from the broker watchlist threshold in B14.",
     color=GREY_TXT, wrap=True)

# ---------------------------------------------------------------- Dim_Fund  BC..BK
DF = [
    ("Client_First_In_TOP",
     '=IF($A{r}="","",IF($U{r}="",0,IF(COUNTIFS($U$2:$U{r},$U{r},$C$2:$C{r},$C{r})=1,1,0)))', "0", 12,
     "1 on the first fund of each client within a TOP. Dim_TOP[Clients_In_TOP] sums it to a distinct client count."),
    ("Broker_Dependent",
     '=IF($A{r}="","",IF(UPPER($P{r})="YES","Yes","No"))', "General", 12,
     "Has_Broker_Process = Yes. The 'Broker (dependent)' supplier type on page 1."),
    ("Thin_Custody_Slack",
     '=IF($A{r}="","",IF($N{r}="No","n/a",IF(AM{r}<=Constants!$B$17,"Yes","No")))', "General", 12,
     "Slack_Custody_Mins at or below Constants!B17 (0 = zero tolerance)."),
    ("Thin_TA_Slack",
     '=IF($A{r}="","",IF($O{r}="No","n/a",IF(AN{r}<=Constants!$B$17,"Yes","No")))', "General", 12,
     "Slack_TA_Mins at or below Constants!B17."),
    ("Thin_Broker_Slack",
     '=IF($A{r}="","",IF($P{r}="No","n/a",IF(AO{r}<=Constants!$B$17,"Yes","No")))', "General", 12,
     "Slack_Broker_Mins at or below Constants!B17. The broker WATCHLIST still uses Constants!B14."),
    ("Custody_Slack_Group",
     '=IF($A{r}="","",IF($N{r}="No","No custody process",IF(BE{r}="Yes","Thin custody slack","Ample custody slack")))',
     "General", 20, "Thin / Ample custody slack group used on page 1."),
    ("Slack_Custody_Num",
     '=IF($A{r}="","",IF($N{r}="No","",AM{r}))', "#,##0", 12,
     "Numeric copy of Slack_Custody_Mins, blank instead of n/a, so Power BI can take a MEDIAN."),
    ("Slack_TA_Num",
     '=IF($A{r}="","",IF($O{r}="No","",AN{r}))', "#,##0", 12, "Numeric copy of Slack_TA_Mins."),
    ("Slack_Broker_Num",
     '=IF($A{r}="","",IF($P{r}="No","",AO{r}))', "#,##0", 12, "Numeric copy of Slack_Broker_Mins."),
]
add_columns(wb["Dim_Fund"], "BC", DF, 201, "v4.5")

# ---------------------------------------------------------------- Fact_Milestone L..M
FM = [
    ("Late_Mins", '=IF($A{r}="","",IF($H{r}="","",MAX(0,$H{r})))', F_NUM, 11,
     "Minutes the file arrived after the binding cut-off, 0 when on time. Page 1 'created' bar."),
    ("Breached_Fund_Days",
     '=IF($A{r}="","",IF($F{r}="","",COUNTIFS(Fact_Work!$B:$B,$A{r},Fact_Work!$AL:$AL,$B{r},Fact_Work!$AI:$AI,"Yes")))',
     "0", 12, "Fund-days of this TOP on this date that breached the client cut-off. Page 1 'reached' bar."),
]
add_columns(wb["Fact_Milestone"], "L", FM, 401, "v4.5")

# ---------------------------------------------------------------- Dim_TOP J..O
DT = [
    ("Clients_In_TOP", '=IF($A{r}="","",COUNTIFS(Dim_Fund!$U:$U,$A{r},Dim_Fund!$BC:$BC,1))', "0", 11,
     "Distinct clients with at least one fund in this TOP."),
    ("Share_Of_Funds",
     '=IF($A{r}="","",IF(COUNTA(Dim_Fund!$A:$A)<=1,"",$D{r}/(COUNTA(Dim_Fund!$A:$A)-1)))', "0%", 11,
     "Funds_In_TOP as a share of all funds. 'Largest TOP' KPI."),
    ("TOP_Days_Scored", '=IF($A{r}="","",COUNTIFS(Fact_Milestone!$B:$B,$A{r},Fact_Milestone!$I:$I,"?*"))', "0", 11,
     "Deliveries scored against a binding cut-off (a fund of the TOP ran that day)."),
    ("TOP_Days_Late", '=IF($A{r}="","",COUNTIFS(Fact_Milestone!$B:$B,$A{r},Fact_Milestone!$I:$I,"No"))', "0", 11,
     "Scored deliveries that missed the binding cut-off."),
    ("Late_Hours_Total", '=IF($A{r}="","",ROUND(SUMIFS(Fact_Milestone!$L:$L,Fact_Milestone!$B:$B,$A{r})/60,2))',
     "0.0", 11, "Sum of Late_Mins / 60. Static whole-period view; Power BI uses [Late Hours] instead."),
    ("Breached_Fund_Days_When_Late",
     '=IF($A{r}="","",COUNTIFS(Fact_Work!$AL:$AL,$A{r},Fact_Work!$AI:$AI,"Yes",Fact_Work!$BB:$BB,"No"))', "0", 14,
     "Breached fund-days on days this TOP was late. Static view; Power BI uses the measure."),
]
add_columns(wb["Dim_TOP"], "J", DT, 200, "v4.5")

# ---------------------------------------------------------------- Fact_Work BA..BQ
FW = [
    ("Client_Name", '=IF($A{r}="","",IFERROR(VLOOKUP($C{r},Dim_Fund!$A:$D,4,0),""))', "General", 20,
     "From Dim_Fund."),
    ("TOP_SLA_Met",
     '=IF($A{r}="","",IF($AL{r}="","",IF(COUNTIFS(Fact_Milestone!$A:$A,$B{r},Fact_Milestone!$B:$B,$AL{r},Fact_Milestone!$I:$I,"No")>0,"No",'
     'IF(COUNTIFS(Fact_Milestone!$A:$A,$B{r},Fact_Milestone!$B:$B,$AL{r},Fact_Milestone!$I:$I,"Yes")>0,"Yes",""))))',
     "General", 11, "Fact_Milestone[SLA_Met] for this fund's TOP on this process date (the supplier measure)."),
    ("Maker_Free_From_DT",
     '=IF($A{r}="","",IF(OR($L{r}="",$W{r}=""),"",IF(COUNTIFS($L:$L,$L{r},$B:$B,$B{r},$W:$W,"<"&$W{r})=0,"",'
     '_xlfn.MAXIFS($W:$W,$L:$L,$L{r},$B:$B,$B{r},$W:$W,"<"&$W{r}))))',
     F_DT, 17, "RECONSTRUCTED. When this maker finished their previous fund on the same process date "
     "(latest Calculated before this fund's). Blank if this was their first fund of the day."),
    ("Queue_Wait_Mins",
     '=IF($A{r}="","",IF($AC{r}="","",IF($BC{r}="",0,MAX(0,MIN($AC{r},ROUND(($BC{r}-$U{r})*1440,0))))))',
     F_NUM, 11, "RECONSTRUCTED. Part of Wait_After_Pricing_Mins spent behind the maker's other funds "
     "(pricing arrival to Maker_Free_From). Purple bar."),
    ("Idle_Wait_Mins", '=IF($A{r}="","",IF($AC{r}="","",$AC{r}-$BD{r}))', F_NUM, 11,
     "RECONSTRUCTED. Rest of the wait: maker free of other funds, so recon or upstream data. Amber bar."),
    ("Broker_Dependent", '=IF($A{r}="","",IFERROR(VLOOKUP($C{r},Dim_Fund!$A:$BD,56,0),""))', "General", 12,
     "From Dim_Fund."),
    ("Custody_Slack_Group", '=IF($A{r}="","",IFERROR(VLOOKUP($C{r},Dim_Fund!$A:$BH,60,0),""))', "General", 20,
     "From Dim_Fund."),
    ("Possible_Cause",
     '=IF($A{r}="","",IF($AI{r}<>"Yes","Not breached",IF($AM{r}="Yes","Pricing late",IF($BG{r}="Thin custody slack",'
     '"Thin custody slack",IF($BF{r}="Yes","Broker-dependent","None of these")))))',
     "General", 18, "Breached fund-days only, first match wins: Pricing late, Thin custody slack, Broker-dependent, "
     "None of these. Only pricing is a measured cause; the others are possibilities."),
    ("Client_Cutoff_Hr", '=IF($A{r}="","",IFERROR(IF(VLOOKUP($C{r},Dim_Fund!$A:$AI,35,0)>0,VLOOKUP($C{r},Dim_Fund!$A:$AI,35,0),""),""))',
     "0.00", 11, "Client cut-off as hours of day. Page 2 red marker. Blank when missing (0)."),
    ("Fund_Label", '=IF($A{r}="","",$C{r}&IF($E{r}="Client","*",""))', "General", 12,
     "Fund_ID with * for client sign-off funds (no Controlled step). Page 2 Y axis."),
    ("Controlled_To_Delivered_Mins", '=IF($A{r}="","",IFERROR(ROUND(($Y{r}-$X{r})*1440,0),""))', F_NUM, 13,
     "NAV_End to Delivered: check plus external tail. For client sign-off funds, FA signed off to delivered."),
    ("Headroom_At_Delivery_Mins", '=IF($A{r}="","",IF($AH{r}="","",-$AH{r}))', F_NUM, 13,
     "Client cut-off minus delivery. Negative = breach. Same number as Breach_Mins with the sign flipped."),
    ("Bar_Offset_Hr", '=IF($A{r}="","",IF($AQ{r}="","",IF($AP{r}="",$AQ{r},MIN($AP{r},$AQ{r}))))', "0.0000", 11,
     "Transparent first segment of the page 2 timeline: pricing arrival, or NAV start if earlier."),
    ("Bar_Wait_Hr", '=IF($A{r}="","",IF($BM{r}="","",ROUND(N($AC{r})/60,4)))', "0.0000", 11,
     "Purple: pricing arrived to NAV started, hours."),
    ("Bar_NAV_Hr", '=IF($A{r}="","",IF(OR($BM{r}="",$AD{r}=""),"",ROUND($AD{r}/60,4)))', "0.0000", 11,
     "Green: NAV started to NAV_End (Controlled; FA signed off for client sign-off funds), hours."),
    ("Bar_Control_Hr", '=IF($A{r}="","",IF($BM{r}="","",ROUND(N($AE{r})/60,4)))', "0.0000", 11,
     "Blue: Controlled to FA signed off, hours. 0 for client sign-off funds."),
    ("Bar_Tail_Hr", '=IF($A{r}="","",IF(OR($BM{r}="",$AF{r}=""),"",ROUND($AF{r}/60,4)))', "0.0000", 11,
     "Amber: FA signed off to delivered, hours. Offset plus the four bars ends at Delivered."),
]
add_columns(wb["Fact_Work"], "BA", FW, 301, "v4.5")
# lookup column indexes must point at the intended Dim_Fund columns
assert cidx("BD") == 56 and cidx("BH") == 60 and cidx("AI") == 35

# ---------------------------------------------------------------- FIX 1: Fact_Work V (missing client cut-off)
ws = wb["Fact_Work"]
for r in range(2, 302):
    old_v = f'=IF($A{r}="","",$B{r}+IFERROR(VLOOKUP($C{r},Dim_Fund!$A:$AI,35,0),0)/24)'
    assert ws[f"V{r}"].value == old_v, (r, ws[f"V{r}"].value)
    ws[f"V{r}"].value = (f'=IF($A{r}="","",IFERROR(IF(VLOOKUP($C{r},Dim_Fund!$A:$AI,35,0)>0,'
                         f'$B{r}+VLOOKUP($C{r},Dim_Fund!$A:$AI,35,0)/24,""),""))')
    # consequence of FIX 1: AA subtracted V without a guard
    old_aa = f'=IF($A{r}="","",IF(U{r}="","",ROUND((V{r}-U{r})*1440,0)))'
    assert ws[f"AA{r}"].value == old_aa, (r, ws[f"AA{r}"].value)
    ws[f"AA{r}"].value = f'=IF($A{r}="","",IF(OR(U{r}="",V{r}=""),"",ROUND((V{r}-U{r})*1440,0)))'

# ---------------------------------------------------------------- FIX 2: Checks C8 text
ws = wb["Checks"]
assert "Constants!B13" in ws["C8"].value
ws["C8"].value = ws["C8"].value.replace("Constants!B13", "Constants!B14")

# ---------------------------------------------------------------- FIX 3: Dim_Signoff notes out of data rows
ws = wb["Dim_Signoff"]
sig_notes = [ws["A6"].value, ws["A7"].value]
for a in ("A6", "A7"):
    ws[a].value = None

# ---------------------------------------------------------------- Static lookup sheets (disconnected / legend order)
def static_sheet(name, headers, rows):
    s = wb.create_sheet(name)
    for j, h in enumerate(headers, 1):
        hdr(s.cell(1, j), h, GREEN)
        s.column_dimensions[get_column_letter(j)].width = 26 if j == 1 else 12
    for i, row in enumerate(rows, 2):
        for j, v in enumerate(row, 1):
            body(s.cell(i, j), v, "General", BLUE_TXT)
    s.freeze_panes = "A2"
    return s

static_sheet("Dim_SupplierType", ["Supplier_Type", "Type_Order"],
             [["Pricing (TOP)", 1], ["Custody", 2], ["TA", 3], ["Broker (dependent)", 4]])
static_sheet("Dim_WaitGroup", ["Wait_Group", "Group_Order"],
             [["Broker-dependent", 1], ["Not broker-dependent", 2], ["Thin custody slack", 3], ["Ample custody slack", 4]])
static_sheet("Dim_Cause", ["Possible_Cause", "Cause_Order"],
             [["Pricing late", 1], ["Thin custody slack", 2], ["Broker-dependent", 3], ["None of these", 4],
              ["Not breached", 5]])
static_sheet("Dim_DetailRow", ["Detail", "Detail_Order"],
             [["Pricing cut-off", 1], ["Pricing arrived", 2], ["NAV started", 3], ["Controlled", 4],
              ["FA signed off", 5], ["Delivered", 6], ["Client cut-off", 7], ["Headroom", 8],
              ["Waited after pricing", 9], ["NAV to Controlled", 10]])

# ---------------------------------------------------------------- Checks additions
ws = wb["Checks"]
r0 = 56
hdr(ws[f"A{r0}"], "POWER BI COLUMNS (v4.5)", wrap=False)
for c in ("B", "C"):
    ws[f"{c}{r0}"].fill = PatternFill("solid", fgColor=NAVY)
checks = [
    ("Fact_Work rows without v4.5 formulas", '=COUNT(Fact_Work!$A:$A)-(COUNTIF(Fact_Work!$BJ:$BJ,"?*")-1)',
     "Target 0. Fill Fact_Work BA to BQ down to the last data row, together with D to AZ."),
    ("Fund-days with no client cut-off", '=COUNTIFS(Fact_Work!$A:$A,"<>",Fact_Work!$V:$V,"")',
     "Target 0. The fund's client cut-off is missing (0) in Dim_Fund, so breach is not scored. Exclude, do not guess."),
    ("Queue + idle not equal to wait", '=COUNTIF(Fact_Work!$BE:$BE,"<0")',
     "Target 0. Idle wait below zero means the queue split is broken for that row."),
    ("Share of wait spent idle", '=IFERROR(SUM(Fact_Work!$BE:$BE)/(SUM(Fact_Work!$BD:$BD)+SUM(Fact_Work!$BE:$BE)),"")',
     "Page 1 'Idle share of wait'. Reconstructed: the maker is assumed busy until their previous fund's Calculated time."),
    ("Funds with thin custody slack", '=COUNTIF(Dim_Fund!$BE:$BE,"Yes")',
     "Slack at or below Constants!B17. With 0 tolerance, only funds whose custody cut-off is at or after pricing."),
    ("Funds with thin TA slack", '=COUNTIF(Dim_Fund!$BF:$BF,"Yes")', "As above for TA."),
    ("Broker-dependent funds", '=COUNTIF(Dim_Fund!$BD:$BD,"Yes")', "Has_Broker_Process = Yes."),
    ("Breached fund-days with a possible cause", '=COUNTIFS(Fact_Work!$BH:$BH,"?*",Fact_Work!$BH:$BH,"<>Not breached")-COUNTIF(Fact_Work!$BH:$BH,"Possible_Cause")',
     "Must equal 'Breached cycles' above."),
]
for i, (a, b, c) in enumerate(checks, r0 + 1):
    body(ws[f"A{i}"], a)
    body(ws[f"B{i}"], b, "0%" if "Share of wait" in a else "General")
    body(ws[f"C{i}"], c, color=GREY_TXT, wrap=True)

# ---------------------------------------------------------------- README additions
ws = wb["README"]
readme = [
    ("VERSION 4.5 — POWER BI PAGES", None, True),
    ("What changed", "Columns added to the right of every existing table so the two report pages (Supplier Concentration "
     "Risk, Fund Level) can be built without Power Query logic. No existing column moved or was renamed. All DAX ships "
     "as a Tabular Editor C# script (NAV_Bottleneck_Measures_v4_5.csx). PBI_Mapping lists every visual, field and measure.", False),
    ("New sheets", "Dim_SupplierType, Dim_WaitGroup, Dim_DetailRow: disconnected row lists for SWITCH measures. "
     "Dim_Cause: legend order for Possible_Cause, related to Fact_Work. Keep all four free of notes: they load as data. "
     "PBI_Mapping is documentation: do not load it.", False),
    ("Thin slack", "Constants!B17, default 0 (zero tolerance). A feed is thin when its slack to the pricing cut-off is at "
     "or below it, so a late file reaches NAV directly. The broker watchlist keeps its own threshold in B14.", False),
    ("Queue vs idle wait", "Wait_After_Pricing_Mins is split in two. Queue = pricing arrival until the maker finished their "
     "previous fund that day (Maker_Free_From_DT, latest Calculated before this fund's). Idle = the rest: the maker was "
     "free of other funds, so the wait points to recon or upstream data. RECONSTRUCTED, not observed: it assumes the maker "
     "was busy continuously until that previous fund, which overstates queue when they had gaps.", False),
    ("Possible cause", "Breached fund-days only, first match wins: Pricing late (this fund's own PSA cut-off), Thin custody "
     "slack, Broker-dependent, None of these. Custody, TA and broker arrivals are not captured, so only pricing is a "
     "measured cause.", False),
    ("Timeline bars", "Bar_Offset_Hr (transparent) + Bar_Wait_Hr + Bar_NAV_Hr + Bar_Control_Hr + Bar_Tail_Hr = delivery "
     "time of day. Client sign-off funds have no Controlled step: their NAV bar runs to FA signed off and the control bar is 0.", False),
    ("Fix: missing client cut-off", "Fact_Work[Client_Cutoff_DT] used to fall back to midnight when the client cut-off was "
     "missing, inventing a breach. It is now blank, and Headroom_After_Pricing_Mins is guarded for it. Checks counts them.", False),
    ("Fix: Checks C8", "Pointed to Constants!B13 (business days); the broker threshold is B14.", False),
    ("Moved from Dim_Signoff", sig_notes[0] + " " + sig_notes[1], False),
    ("Filling down", "Fact_Work formulas now run D to AZ and BA to BQ; Fact_Milestone D to M; Dim_Fund to BK; Dim_TOP to O.", False),
    ("Power Query", PQ_NOTE, False),
]
r = 107
for a, b, is_hdr in readme:
    if is_hdr:
        hdr(ws[f"A{r}"], a, wrap=False)
        ws[f"B{r}"].fill = PatternFill("solid", fgColor=NAVY)
    else:
        body(ws[f"A{r}"], a)
        body(ws[f"B{r}"], b, wrap=True)
    r += 1
body(ws["B25"], ws["B25"].value + " v4.5 adds 9 derived columns (BC to BK).")
body(ws["B28"], ws["B28"].value + " v4.5 adds 17 Power BI columns (BA to BQ).")
ws["B25"].alignment = ws["B28"].alignment = Alignment(wrap_text=True)

# ---------------------------------------------------------------- PBI_Mapping sheet
pm = wb.create_sheet("PBI_Mapping")
pm.sheet_properties.tabColor = "123A5E"
row = 1
hdr(pm.cell(row, 1), "POWER BI MAPPING — documentation only, do not load this sheet", wrap=False)
for c in range(2, 8):
    pm.cell(row, c).fill = PatternFill("solid", fgColor=NAVY)
row += 1
body(pm.cell(row, 1), "Run NAV_Bottleneck_Measures_v4_5.csx in Tabular Editor (C# Script tab) against the model loaded "
     "from this workbook. It creates the relationships, sort orders and every measure below, and is safe to re-run.",
     wrap=True)
pm.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
pm.row_dimensions[row].height = 30
row += 2
hdr(pm.cell(row, 1), "VISUALS", wrap=False)
for c in range(2, 8):
    pm.cell(row, c).fill = PatternFill("solid", fgColor=NAVY)
row += 1
for j, h in enumerate(["Page", "Visual", "Type", "Axis / rows", "Values / measures", "Legend / filters / sort", "Notes"], 1):
    hdr(pm.cell(row, j), h, GREEN)
row += 1
for v in VISUALS:
    for j, val in enumerate(v, 1):
        body(pm.cell(row, j), val, wrap=True)
    row += 1
row += 1
hdr(pm.cell(row, 1), "RELATIONSHIPS (created by the script)", wrap=False)
for c in range(2, 8):
    pm.cell(row, c).fill = PatternFill("solid", fgColor=NAVY)
row += 1
for j, h in enumerate(["From (many)", "To (one)", "Notes"], 1):
    hdr(pm.cell(row, j), h, GREEN)
row += 1
for ft, fc, tt, tc, note in RELATIONSHIPS:
    body(pm.cell(row, 1), f"{ft}[{fc}]")
    body(pm.cell(row, 2), f"{tt}[{tc}]")
    body(pm.cell(row, 3), note, wrap=True)
    row += 1
row += 1
hdr(pm.cell(row, 1), "MEASURES (created by the script)", wrap=False)
for c in range(2, 8):
    pm.cell(row, c).fill = PatternFill("solid", fgColor=NAVY)
row += 1
for j, h in enumerate(["Measure", "Home table", "Display folder", "Format", "DAX", "Description"], 1):
    hdr(pm.cell(row, j), h, GREEN)
row += 1
for m in MEASURES:
    vals = [m["name"], m["table"], m["folder"], m.get("fmt", ""), m["dax"].strip(), m.get("desc", "")]
    for j, val in enumerate(vals, 1):
        body(pm.cell(row, j), val, wrap=True)
    row += 1
for col, w in zip("ABCDEFG", [30, 30, 22, 30, 60, 40, 50]):
    pm.column_dimensions[col].width = w

wb.save(dst)

# ---------------------------------------------------------------- Tabular Editor C# script
def cs(s):
    return '@"' + s.strip().replace('"', '""') + '"'

lines = []
w = lines.append
w("// NAV Bottleneck — Power BI measures for template v4.5")
w("// Tabular Editor 2 or 3: open the model (or connect to Power BI Desktop), paste into the C# Script tab, run (F5),")
w("// then save (Ctrl+S). Safe to re-run: existing measures are updated in place, existing relationships kept.")
w("// Assumes each workbook sheet was loaded as a table with the same name and column headers.")
w("")
w("var log = new System.Text.StringBuilder();")
w("")
w("Func<string, Table> T = (name) => Model.Tables.FirstOrDefault(t => t.Name == name);")
w("Func<string, string, Column> C = (table, col) => { var t = T(table); return t == null ? null : t.Columns.FirstOrDefault(c => c.Name == col); };")
w("")
w("Action<string, string, string, string, string, string> AddMeasure = (table, name, folder, fmt, desc, dax) => {")
w("    var t = T(table);")
w("    if (t == null) { log.AppendLine(\"SKIPPED measure \" + name + \": table \" + table + \" not in model\"); return; }")
w("    var m = Model.AllMeasures.FirstOrDefault(x => x.Name == name);")
w("    if (m != null && m.Table != t) { m.Delete(); m = null; }")
w("    if (m == null) m = t.AddMeasure(name);")
w("    m.Expression = dax;")
w("    m.DisplayFolder = folder;")
w("    m.Description = desc;")
w("    if (!string.IsNullOrEmpty(fmt)) m.FormatString = fmt;")
w("};")
w("")
w("Action<string, string, string, string> Rel = (ft, fc, tt, tc) => {")
w("    var from = C(ft, fc); var to = C(tt, tc);")
w("    if (from == null || to == null) { log.AppendLine(\"SKIPPED relationship \" + ft + \"[\" + fc + \"] -> \" + tt + \"[\" + tc + \"]: column not in model\"); return; }")
w("    if (Model.Relationships.Any(r => r.FromColumn == from && r.ToColumn == to)) return;")
w("    var rel = Model.AddRelationship();")
w("    rel.FromColumn = from; rel.ToColumn = to;")
w("    rel.FromCardinality = RelationshipEndCardinality.Many; rel.ToCardinality = RelationshipEndCardinality.One;")
w("    rel.CrossFilteringBehavior = CrossFilteringBehavior.OneDirection;")
w("    rel.IsActive = true;")
w("};")
w("")
w("Action<string, string, string> SortBy = (table, col, by) => {")
w("    var c = C(table, col); var b = C(table, by);")
w("    if (c == null || b == null) { log.AppendLine(\"SKIPPED sort \" + table + \"[\" + col + \"]\"); return; }")
w("    c.SortByColumn = b;")
w("};")
w("")
w("// ---- Relationships")
for ft, fc, tt, tc, note in RELATIONSHIPS:
    w(f'Rel("{ft}", "{fc}", "{tt}", "{tc}");  // {note}')
w("")
w("// ---- Sort orders")
for t, c, b in SORT_BY:
    w(f'SortBy("{t}", "{c}", "{b}");')
w("")
w("// ---- Measures")
for m in MEASURES:
    w(f'AddMeasure({cs(m["table"])}, {cs(m["name"])}, {cs(m["folder"])}, {cs(m.get("fmt", ""))}, {cs(m.get("desc", ""))},')
    w(f'    {cs(m["dax"])});')
w("")
w(f'log.Insert(0, "NAV Bottleneck v4.5: {len(MEASURES)} measures processed.\\n");')
w("Info(log.ToString());")
with open(csx, "w", encoding="utf-8") as fh:
    fh.write("\n".join(lines) + "\n")
print("built", dst, csx, len(MEASURES), "measures")
